from pptx import Presentation

from src.ppt_builder import (
    remove_shape_if_exists,
    remove_template_instructions,
    replace_icon,
    replace_placeholder_with_image,
    set_bullets,
    set_text,
)
from src.analysis import calculate_kpis
from src.chart_generation import (
    create_customer_chart,
    create_performance_chart,
    create_risk_chart,
)
from src.config import (
    DATASET_PATH,
    ICON_MAPPING,
    OUTPUT_DIR,
    OUTPUT_PPTX,
    TEMPLATE_PATH,
)
from src.data_loader import load_data
from src.data_validator import validate_data
from src.ppt_builder import (
    replace_icon,
    replace_placeholder_with_image,
    set_bullets,
    set_text,
)


def main():

    print("\n==============================================")
    print("NovaRetail QBR Generator Started")
    print("==============================================\n")

    # ==========================================================
    # 1. CREATE OUTPUT DIRECTORY
    # ==========================================================

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    print(f"Output directory:\n{OUTPUT_DIR}\n")

    # ==========================================================
    # 2. CHECK INPUT FILES
    # ==========================================================

    print("Checking required input files...")

    if not DATASET_PATH.exists():
        raise FileNotFoundError(
            f"\nDataset not found:\n{DATASET_PATH}"
        )

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"\nPowerPoint template not found:\n{TEMPLATE_PATH}"
        )

    print("Dataset found.")
    print("PowerPoint template found.\n")

    # ==========================================================
    # 3. LOAD DATA
    # ==========================================================

    print("Loading NovaRetail dataset...")

    dataframes = load_data(DATASET_PATH)

    print("Dataset loaded successfully.\n")

    # ==========================================================
    # 4. VALIDATE DATA
    # ==========================================================

    print("Validating dataset...")

    validate_data(dataframes)

    print("Dataset structure is valid.\n")

    # ==========================================================
    # 5. GET DATAFRAMES
    # ==========================================================

    performance = dataframes["Regional_Performance"]
    customer = dataframes["Customer_Experience"]
    risk = dataframes["Operational_Risk"]

    # ==========================================================
    # 6. CALCULATE KPIs
    # ==========================================================

    print("Calculating business KPIs...")

    kpis = calculate_kpis(dataframes)

    print("KPIs calculated successfully.\n")

    # ==========================================================
    # 7. GENERATE CHARTS
    # ==========================================================

    print("Generating charts...")

    performance_chart = create_performance_chart(
        performance
    )

    customer_chart = create_customer_chart(
        customer
    )

    risk_chart = create_risk_chart(
        risk
    )

    print("Charts generated successfully.\n")

    # ==========================================================
    # 8. LOAD POWERPOINT TEMPLATE
    # ==========================================================

    print("Loading PowerPoint template...")

    presentation = Presentation(
        TEMPLATE_PATH
    )

    slide_count = len(presentation.slides)

    print(
        f"Template loaded successfully. "
        f"Slides found: {slide_count}\n"
    )

    if slide_count < 5:
        raise ValueError(
            "PowerPoint template must contain at least 5 slides."
        )

    # ==========================================================
    # 9. GET SLIDES
    # ==========================================================

    slide_1 = presentation.slides[0]
    slide_2 = presentation.slides[1]
    slide_3 = presentation.slides[2]
    slide_4 = presentation.slides[3]
    slide_5 = presentation.slides[4]

    # ==========================================================
    # SLIDE 1: EXECUTIVE SUMMARY
    # ==========================================================

    print("Building Slide 1...")

    set_text(
        slide_1,
        "subtitle_1",
        "Q2 2026 | Quarterly Business Review",
    )

    set_text(
        slide_1,
        "kpi1_value",
        f"${kpis['total_revenue_usd_m'] / 1000:.2f}B",
    )

    set_text(
        slide_1,
        "kpi1_label",
        "Total Revenue",
    )

    set_text(
        slide_1,
        "kpi2_value",
        f"{kpis['qoq_revenue_growth_pct']:+.1f}%",
    )

    set_text(
        slide_1,
        "kpi2_label",
        "QoQ Revenue Growth",
    )

    set_text(
        slide_1,
        "kpi3_value",
        f"{kpis['weighted_nps']:.1f}",
    )

    set_text(
        slide_1,
        "kpi3_label",
        "Sales-Weighted NPS",
    )

    # Correct Slide 1 template shape: summary_text
    set_text(
        slide_1,
        "summary_text",
        (
            "NovaRetail generated $1.28B in Q2 revenue, representing "
            "4.5% sequential growth. APAC delivered the strongest "
            "growth momentum and maintained the highest customer "
            "advocacy, while North America remained the largest "
            "revenue contributor. Latin America continued to contract "
            "as customer experience and operational risk deteriorated. "
            "APAC capacity also requires action to protect future growth."
        ),
    )

    replace_icon(
        slide_1,
        "icon_1",
        ICON_MAPPING[1],
    )

    print("Slide 1 completed.\n")

    # ==========================================================
    # SLIDE 2: REGIONAL PERFORMANCE
    # ==========================================================

    print("Building Slide 2...")

    set_text(
        slide_2,
        "title_2",
        (
            "APAC Leads Growth as Latin America "
            "Revenue Falls Further"
        ),
    )

    replace_placeholder_with_image(
        slide_2,
        "chart_performance",
        performance_chart,
    )
    remove_template_instructions(slide_2)
    remove_shape_if_exists(
    slide_2,
    "EXACT_SLIDE_2_INSTRUCTION_SHAPE_NAME",
)

    set_bullets(
        slide_2,
        "insights_performance",
        [
            "Total Q2 revenue reached $1.28B, up 4.5% QoQ.",
            "APAC revenue increased from $260M to $330M.",
            "North America remained the largest region at $465M.",
            (
                "Latin America fell to $140M "
                "as returns increased to 8.2%."
            ),
        ],
    )

    replace_icon(
        slide_2,
        "icon_2",
        ICON_MAPPING[2],
    )
    remove_shape_if_exists(
    slide_2,
    "icon_2_label",
)

    print("Slide 2 completed.\n")

    # ==========================================================
    # SLIDE 3: CUSTOMER EXPERIENCE
    # ==========================================================

    print("Building Slide 3...")

    set_text(
        slide_3,
        "title_3",
        (
            "Latin America Customer Experience Deteriorates "
            "as APAC Extends Its Lead"
        ),
    )

    replace_placeholder_with_image(
        slide_3,
        "chart_customer",
        customer_chart,
    )
    remove_template_instructions(slide_3)
    remove_shape_if_exists(
    slide_3,
    "EXACT_SLIDE_3_INSTRUCTION_SHAPE_NAME",
)

    set_bullets(
        slide_3,
        "insights_customer",
        [
            "APAC achieved the highest Q2 NPS at 62.",
            "APAC delivery performance improved to 2.2 days.",
            "Latin America NPS fell from 38 to 29.",
            (
                "Latin America tickets climbed to 31K "
                "as delivery reached 7.5 days."
            ),
        ],
    )

    replace_icon(
        slide_3,
        "icon_3",
        ICON_MAPPING[3],
    )
    remove_shape_if_exists(
    slide_3,
    "icon_3_label",
)

    print("Slide 3 completed.\n")

    # ==========================================================
    # SLIDE 4: OPERATIONAL RISK
    # ==========================================================

    print("Building Slide 4...")

    set_text(
        slide_4,
        "title_4",
        (
            "Latin America Risk Escalates While "
            "APAC Nears Capacity Limits"
        ),
    )

    replace_placeholder_with_image(
        slide_4,
        "chart_risk",
        risk_chart,
    )
    remove_template_instructions(slide_4)
    remove_shape_if_exists(
    slide_4,
    "EXACT_SLIDE_4_INSTRUCTION_SHAPE_NAME",
)

    set_bullets(
        slide_4,
        "insights_risk",
        [
            "Latin America stockouts reached 12.5% in Q2.",
            (
                "Fraud incidents in Latin America "
                "increased from 18 to 30."
            ),
            (
                "Latin America return processing "
                "increased to 10 days."
            ),
            (
                "APAC warehouse utilization reached "
                "a critical 92%."
            ),
        ],
    )

    replace_icon(
        slide_4,
        "icon_4",
        ICON_MAPPING[4],
    )
    remove_shape_if_exists(
    slide_4,
    "icon_4_label",
)

    print("Slide 4 completed.\n")

    # ==========================================================
    # SLIDE 5: RECOMMENDED ACTIONS
    # ==========================================================

    print("Building Slide 5...")

    set_text(
        slide_5,
        "title_5",
        "Recommended Actions and Executive Priorities",
    )

    # Correct Slide 5 template shape: recommendations_text
    set_bullets(
        slide_5,
        "recommendations_text",
        [
            (
                "Stabilize Latin America operations through "
                "a cross-functional recovery plan."
            ),
            (
                "Expand APAC warehouse and fulfillment "
                "capacity ahead of future demand."
            ),
            (
                "Replicate APAC delivery practices across "
                "lower-performing regions."
            ),
            (
                "Introduce monthly alerts for stockouts, "
                "NPS, fraud and capacity utilization."
            ),
        ],
    )

    set_text(
        slide_5,
        "closing_statement",
        (
            "Prioritize immediate operational recovery in "
            "Latin America while investing ahead of "
            "APAC's continued growth."
        ),
    )

    replace_icon(
        slide_5,
        "icon_5",
        ICON_MAPPING[5],
    )

    print("Slide 5 completed.\n")

    # ==========================================================
    # 10. SAVE PRESENTATION
    # ==========================================================

    print("Saving presentation...")

    presentation.save(
        OUTPUT_PPTX
    )

    # ==========================================================
    # 11. VERIFY GENERATED OUTPUT
    # ==========================================================

    if not OUTPUT_PPTX.exists():
        raise FileNotFoundError(
            "\nPowerPoint generation completed, but "
            "the output file was not found.\n"
            f"Expected location:\n{OUTPUT_PPTX}"
        )

    file_size_mb = (
        OUTPUT_PPTX.stat().st_size
        / (1024 * 1024)
    )

    print("\n==============================================")
    print("NovaRetail presentation created successfully!")
    print("==============================================")

    print(
        f"\nOutput file:\n{OUTPUT_PPTX}"
    )

    print(
        f"\nFile size: {file_size_mb:.2f} MB"
    )

    print("\n==============================================\n")


# ==============================================================
# APPLICATION ENTRY POINT
# ==============================================================

if __name__ == "__main__":
    main()