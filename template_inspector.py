from pathlib import Path
from pptx import Presentation


BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = BASE_DIR / "assets" / "novaretail_template.pptx"



def inspect_template(template_path: Path) -> None:
    presentation = Presentation(template_path)

    print(f"Total slides: {len(presentation.slides)}")
    print(
        f"Slide size: "
        f"{presentation.slide_width / 914400:.2f} x "
        f"{presentation.slide_height / 914400:.2f} inches"
    )

    for slide_number, slide in enumerate(presentation.slides, start=1):
        print("\n" + "=" * 80)
        print(f"SLIDE {slide_number}")
        print("=" * 80)

        for shape in slide.shapes:
            text = ""

            if hasattr(shape, "text"):
                text = shape.text.replace("\n", " | ")[:100]

            print(
                f"Name={shape.name!r}, "
                f"Type={shape.shape_type}, "
                f"Left={shape.left / 914400:.2f}, "
                f"Top={shape.top / 914400:.2f}, "
                f"Width={shape.width / 914400:.2f}, "
                f"Height={shape.height / 914400:.2f}, "
                f"Text={text!r}"
            )


if __name__ == "__main__":
    inspect_template(TEMPLATE_PATH)