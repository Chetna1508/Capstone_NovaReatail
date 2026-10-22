from pptx import Presentation


FORBIDDEN_TEXT = [
    "[FILL:",
    "[VALUE]",
    "Insert your chart here",
    "Pick the icon",
]


def check_instruction_text(pptx_path):
    presentation = Presentation(pptx_path)
    issues = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            for forbidden in FORBIDDEN_TEXT:
                if forbidden.lower() in text.lower():
                    issues.append(
                        {
                            "slide": slide_number,
                            "shape": shape.name,
                            "text": text,
                        }
                    )

    if issues:
        raise ValueError(
            f"Instruction text was found: {issues}"
        )

    print("No template instruction text remains.")