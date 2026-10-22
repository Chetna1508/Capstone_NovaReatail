from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt


# ============================================================
# COLORS
# ============================================================

WHITE = RGBColor(255, 255, 255)

LIGHT_TEXT = RGBColor(230, 234, 245)

MUTED_TEXT = RGBColor(190, 200, 225)

DARK_TEXT = RGBColor(25, 31, 55)


# ============================================================
# FIND SHAPE
# ============================================================

def find_shape(slide, shape_name):

    for shape in slide.shapes:

        if shape.name == shape_name:
            return shape

    available_names = [
        shape.name
        for shape in slide.shapes
    ]

    raise KeyError(
        f"Shape {shape_name!r} was not found. "
        f"Available shapes: {available_names}"
    )


# ============================================================
# REMOVE SHAPE
# ============================================================

def remove_shape(shape):

    shape_element = shape._element

    shape_element.getparent().remove(
        shape_element
    )


# ============================================================
# REPLACE PLACEHOLDER WITH IMAGE
# ============================================================

def replace_placeholder_with_image(
    slide,
    shape_name,
    image_path,
):

    placeholder = find_shape(
        slide,
        shape_name
    )

    image_path = Path(image_path)

    if not image_path.exists():

        raise FileNotFoundError(
            f"Image was not found:\n{image_path}"
        )

    left = placeholder.left
    top = placeholder.top
    width = placeholder.width
    height = placeholder.height

    remove_shape(
        placeholder
    )

    picture = slide.shapes.add_picture(
        str(image_path),
        left,
        top,
        width,
        height,
    )

    picture.name = (
        f"{shape_name}_generated"
    )

    return picture


# ============================================================
# REPLACE ICON
# ============================================================

def replace_icon(
    slide,
    shape_name,
    icon_path,
):

    return replace_placeholder_with_image(
        slide=slide,
        shape_name=shape_name,
        image_path=icon_path,
    )


# ============================================================
# DETERMINE TEXT STYLE
# ============================================================

def get_text_style(shape_name):

    """
    Returns:
        font_size
        bold
        color
        alignment
    """

    # --------------------------------------------------------
    # KPI VALUES
    # --------------------------------------------------------

    if shape_name in {
        "kpi1_value",
        "kpi2_value",
        "kpi3_value",
    }:

        return (
            26,
            True,
            WHITE,
            PP_ALIGN.CENTER,
        )

    # --------------------------------------------------------
    # KPI LABELS
    # --------------------------------------------------------

    if shape_name in {
        "kpi1_label",
        "kpi2_label",
        "kpi3_label",
    }:

        return (
            13,
            False,
            LIGHT_TEXT,
            PP_ALIGN.CENTER,
        )

    # --------------------------------------------------------
    # SLIDE 1 SUBTITLE
    # --------------------------------------------------------

    if shape_name == "subtitle_1":

        return (
            16,
            False,
            LIGHT_TEXT,
            PP_ALIGN.LEFT,
        )

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    if shape_name == "summary_text":

        return (
            15,
            False,
            LIGHT_TEXT,
            PP_ALIGN.LEFT,
        )

    # --------------------------------------------------------
    # SLIDE 5 CLOSING STATEMENT
    # --------------------------------------------------------

    if shape_name == "closing_statement":

        return (
            15,
            True,
            LIGHT_TEXT,
            PP_ALIGN.LEFT,
        )

    # --------------------------------------------------------
    # SLIDE TITLES
    # --------------------------------------------------------

    if shape_name in {
        "title_2",
        "title_3",
        "title_4",
    }:

        return (
            22,
            True,
            DARK_TEXT,
            PP_ALIGN.LEFT,
        )

    # Slide 5 has dark background
    if shape_name == "title_5":

        return (
            22,
            True,
            WHITE,
            PP_ALIGN.LEFT,
        )

    # --------------------------------------------------------
    # DEFAULT
    # --------------------------------------------------------

    return (
        None,
        None,
        None,
        None,
    )


# ============================================================
# SET TEXT
# ============================================================

def set_text(
    slide,
    shape_name,
    text,
    font_size=None,
    bold=None,
    color=None,
    alignment=None,
):

    shape = find_shape(
        slide,
        shape_name
    )

    if not shape.has_text_frame:

        raise ValueError(
            f"{shape_name!r} "
            "does not contain a text frame."
        )

    text_frame = shape.text_frame

    text_frame.clear()

    text_frame.word_wrap = True

    # --------------------------------------------------------
    # Obtain automatic styling
    # --------------------------------------------------------

    (
        default_size,
        default_bold,
        default_color,
        default_alignment,
    ) = get_text_style(
        shape_name
    )

    if font_size is None:
        font_size = default_size

    if bold is None:
        bold = default_bold

    if color is None:
        color = default_color

    if alignment is None:
        alignment = default_alignment

    # --------------------------------------------------------
    # Add text
    # --------------------------------------------------------

    paragraph = text_frame.paragraphs[0]

    paragraph.text = str(text)

    if alignment is not None:
        paragraph.alignment = alignment

    # --------------------------------------------------------
    # Format runs
    # --------------------------------------------------------

    for run in paragraph.runs:

        if font_size is not None:
            run.font.size = Pt(
                font_size
            )

        if bold is not None:
            run.font.bold = bold

        if color is not None:

            if isinstance(
                color,
                RGBColor,
            ):

                run.font.color.rgb = color

            else:

                run.font.color.rgb = RGBColor(
                    *color
                )

    return shape


# ============================================================
# SET BULLETS
# ============================================================

def set_bullets(
    slide,
    shape_name,
    bullets,
    font_size=14,
    color=None,
):

    shape = find_shape(
        slide,
        shape_name
    )

    if not shape.has_text_frame:

        raise ValueError(
            f"{shape_name!r} "
            "does not have a text frame."
        )

    text_frame = shape.text_frame

    text_frame.clear()

    text_frame.word_wrap = True

    # --------------------------------------------------------
    # Determine appropriate color
    # --------------------------------------------------------

    if color is None:

        if shape_name == "recommendations_text":

            text_color = LIGHT_TEXT

        else:

            text_color = DARK_TEXT

    else:

        if isinstance(
            color,
            RGBColor,
        ):

            text_color = color

        else:

            text_color = RGBColor(
                *color
            )

    # --------------------------------------------------------
    # Add bullet paragraphs
    # --------------------------------------------------------

    for index, bullet in enumerate(
        bullets
    ):

        if index == 0:

            paragraph = (
                text_frame.paragraphs[0]
            )

        else:

            paragraph = (
                text_frame.add_paragraph()
            )

        paragraph.text = str(bullet)

        paragraph.level = 0

        paragraph.space_after = Pt(8)

        paragraph.line_spacing = 1.05

        # Add bullet character ourselves.
        # This gives consistent output across placeholders.

        if not paragraph.text.startswith("•"):

            paragraph.text = (
                f"• {paragraph.text}"
            )

        for run in paragraph.runs:

            run.font.size = Pt(
                font_size
            )

            run.font.color.rgb = (
                text_color
            )


# ============================================================
# OPTIONAL: FORMAT STANDARD TEXT BOX
# ============================================================

def format_text_box(
    slide,
    shape_name,
    font_size=14,
    bold=False,
    color=(255, 255, 255),
):

    shape = find_shape(
        slide,
        shape_name
    )

    if not shape.has_text_frame:

        return shape

    text_frame = shape.text_frame

    text_frame.word_wrap = True

    for paragraph in text_frame.paragraphs:

        for run in paragraph.runs:

            run.font.size = Pt(
                font_size
            )

            run.font.bold = bold

            run.font.color.rgb = (
                RGBColor(*color)
            )

    return shape
def remove_shape_if_exists(slide, shape_name):
    for shape in slide.shapes:
        if shape.name == shape_name:
            remove_shape(shape)
            return True

    return False
def remove_template_instructions(slide):
    """
    Remove instructional text boxes left in the PowerPoint template,
    such as:
    - CHART
    - Insert your chart here...
    - Suggested: ...
    """

    shapes_to_remove = []

    for shape in slide.shapes:

        if not shape.has_text_frame:
            continue

        text = shape.text.strip()

        if not text:
            continue

        text_lower = text.lower()

        if (
            text_lower == "chart"
            or "insert your chart here" in text_lower
            or "suggested:" in text_lower
        ):
            shapes_to_remove.append(shape)

    for shape in shapes_to_remove:

        print(
            f"Removing template instruction: "
            f"{shape.name!r} -> {shape.text!r}"
        )

        remove_shape(shape)